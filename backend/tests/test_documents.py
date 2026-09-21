"""Tests for Module 4 - Documents, OCR pipeline and AI extraction (mock)."""
from __future__ import annotations

import io

API = "/api/v1"

SAMPLE_LAB_TEXT = (
    "Buletin de analize\n"
    "Glicemie 105 mg/dL (70 - 99)\n"
    "Hemoglobina: 13.5 g/dL 12 - 16\n"
    "Colesterol total 260 mg/dL (0 - 200)\n"
    "Diagnostic: sindrom metabolic\n"
)


def _auth_headers(client, email="doc@example.com"):
    client.post(f"{API}/auth/register", json={"email": email, "password": "Parola1234"})
    tokens = client.post(
        f"{API}/auth/login", json={"email": email, "password": "Parola1234"}
    ).json()
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def _upload(client, headers, content=SAMPLE_LAB_TEXT, filename="analize.pdf",
            content_type="application/pdf", category="lab"):
    return client.post(
        f"{API}/documents",
        headers=headers,
        files={"file": (filename, io.BytesIO(content.encode()), content_type)},
        data={"category": category},
    )


def test_upload_rejects_unsupported_type(client):
    h = _auth_headers(client)
    r = _upload(client, h, filename="note.txt", content_type="text/plain")
    assert r.status_code == 415


def test_upload_rejects_empty_file(client):
    h = _auth_headers(client)
    r = _upload(client, h, content="")
    assert r.status_code == 400


def test_upload_processes_and_extracts_lab_values(client):
    """A PDF with embedded text is parsed by pypdf; the mock AI extracts labs."""
    h = _auth_headers(client)
    # Build a real one-page PDF containing the sample text so pypdf can read it.
    pdf_bytes = _make_text_pdf(SAMPLE_LAB_TEXT)
    r = client.post(
        f"{API}/documents",
        headers=h,
        files={"file": ("analize.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        data={"category": "lab"},
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["status"] == "done"
    analytes = {lr["analyte"]: lr for lr in body["lab_results"]}
    assert "Glicemie" in analytes
    assert analytes["Glicemie"]["value"] == 105.0
    assert analytes["Glicemie"]["flag"] == "high"  # 105 > 99, within 0.5*span
    # 260 vs (0-200): 0.5*span = 100, 260 < 300 -> high (not critical)
    assert analytes["Colesterol total"]["flag"] == "high"


def test_list_filter_get_and_delete(client):
    h = _auth_headers(client)
    up = _upload(client, h)
    assert up.status_code == 201
    doc_id = up.json()["id"]

    # List + category filter
    assert len(client.get(f"{API}/documents", headers=h).json()) == 1
    assert len(client.get(f"{API}/documents?category=lab", headers=h).json()) == 1
    assert client.get(f"{API}/documents?category=ct", headers=h).json() == []

    # Detail
    assert client.get(f"{API}/documents/{doc_id}", headers=h).json()["id"] == doc_id

    # Download presigned URL
    dl = client.get(f"{API}/documents/{doc_id}/download", headers=h).json()
    assert dl["url"].startswith("https://storage.test/")

    # Delete
    assert client.delete(f"{API}/documents/{doc_id}", headers=h).status_code == 204
    assert client.get(f"{API}/documents/{doc_id}", headers=h).status_code == 404


def test_documents_are_isolated_per_patient(client):
    h1 = _auth_headers(client, "p1@example.com")
    h2 = _auth_headers(client, "p2@example.com")
    doc_id = _upload(client, h1).json()["id"]
    # Patient 2 cannot see or fetch patient 1's document.
    assert client.get(f"{API}/documents", headers=h2).json() == []
    assert client.get(f"{API}/documents/{doc_id}", headers=h2).status_code == 404


def _make_text_pdf(text: str) -> bytes:
    """Minimal single-page PDF with a text stream (readable by pypdf)."""
    from pypdf import PdfWriter
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        FloatObject,
        NameObject,
        NumberObject,
    )

    lines = text.splitlines()
    content = "BT /F1 12 Tf 40 800 Td 14 TL " + " ".join(
        f"({line.replace('(', '').replace(')', '')}) Tj T*" for line in lines
    ) + " ET"

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    page = writer.pages[0]

    stream = DecodedStreamObject()
    stream.set_data(content.encode("latin-1"))
    stream_ref = writer._add_object(stream)
    page[NameObject("/Contents")] = stream_ref

    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font)
    resources = DictionaryObject()
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font_ref
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources
    page[NameObject("/MediaBox")] = ArrayObject(
        [NumberObject(0), NumberObject(0), FloatObject(595), FloatObject(842)]
    )

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_original_endpoint_checks_owner(client):
    owner = _auth_headers(client, "owner-original@example.com")
    other = _auth_headers(client, "other-original@example.com")
    data = _make_text_pdf(SAMPLE_LAB_TEXT)
    result = client.post(API + "/documents", headers=owner,
                         files={"file": ("original.pdf", data, "application/pdf")})
    assert result.status_code == 201
    url = API + f"/documents/{result.json()['id']}/original"
    assert client.get(url).status_code == 401
    assert client.get(url, headers=other).status_code == 404
    downloaded = client.get(url, headers=owner)
    assert downloaded.status_code == 200
    assert downloaded.content == data
    assert downloaded.headers["cache-control"] == "no-store"


def test_upload_limit_reads_only_bounded_bytes(client, monkeypatch):
    monkeypatch.setattr("app.api.routes.documents.MAX_SIZE_BYTES", 10)
    headers = _auth_headers(client, "limit-original@example.com")
    assert client.post(API + "/documents", headers=headers,
                       files={"file": ("big.pdf", b"x" * 11, "application/pdf")}
                       ).status_code == 413


def test_docx_paragraphs_and_tables_are_extracted(client):
    from docx import Document as WordDocument

    from app.services.ocr import DOCX_TYPE

    source = WordDocument()
    source.add_paragraph("Buletin fictiv de laborator")
    table = source.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Glicemie", "105 mg/dL", "70 - 99"], strict=True):
        cell.text = value
    buffer = io.BytesIO()
    source.save(buffer)
    data = buffer.getvalue()
    headers = _auth_headers(client)
    result = client.post(API + "/documents", headers=headers,
                         files={"file": ("fictiv.docx", data, DOCX_TYPE)})
    assert result.status_code == 201, result.text
    body = result.json()
    assert body["status"] == "done"
    assert body["category"] == "lab"
    assert "Buletin fictiv" in body["extracted_text"]
    assert body["lab_results"][0]["value"] == 105
    original = client.get(API + f"/documents/{body['id']}/original", headers=headers)
    assert original.content == data


def test_docx_invalid_and_excessively_compressed_rejected(client):
    import zipfile

    from app.services.ocr import DOCX_TYPE

    headers = _auth_headers(client)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "content")
        archive.writestr("word/document.xml", "x" * 200_000)
    for data in [b"not a docx", buffer.getvalue()]:
        result = client.post(API + "/documents", headers=headers,
                             files={"file": ("invalid.docx", data, DOCX_TYPE)})
        assert result.status_code == 400
    assert client.get(API + "/documents", headers=headers).json() == []


def test_empty_extraction_is_failed_and_does_not_invoke_ai(client, monkeypatch):
    from app.services.ai.mock import MockProvider

    calls = []

    def unexpected(*args):
        calls.append(True)
        raise AssertionError("AI must not receive empty extraction")

    monkeypatch.setattr("app.services.ocr.extract_text", lambda *args: "")
    monkeypatch.setattr(MockProvider, "extract_document", unexpected)
    headers = _auth_headers(client)
    result = _upload(client, headers).json()
    assert result["status"] == "failed"
    assert calls == []
    assert result["lab_results"] == []
    assert "Originalul este păstrat" in result["ai_summary"]
    assert client.get(API + f"/documents/{result['id']}/original",
                      headers=headers).status_code == 200


def test_reprocess_failure_rolls_back_replacement_and_hides_exception(client, monkeypatch, caplog):
    headers = _auth_headers(client)
    result = client.post(API + "/documents", headers=headers,
                         files={"file": ("labs.pdf", _make_text_pdf(SAMPLE_LAB_TEXT),
                                         "application/pdf")}).json()
    url = API + f"/documents/{result['id']}"
    previous = result["lab_results"]
    assert len(previous) == 3

    def fail(*args):
        raise RuntimeError("sensitive-provider-content")

    monkeypatch.setattr("app.services.document_processing._persist_lab_values", fail)
    failed = client.post(url + "/reprocess", headers=headers)
    assert failed.status_code == 200
    assert failed.json()["status"] == "failed"
    assert failed.json()["lab_results"] == previous
    assert "sensitive-provider-content" not in failed.text
    assert "sensitive-provider-content" not in caplog.text


def test_reprocessing_does_not_duplicate_results_or_allow_busy_document(client, db_session):
    from app.models.document import Document
    from app.models.enums import ProcessingStatus

    headers = _auth_headers(client)
    result = client.post(API + "/documents", headers=headers,
                         files={"file": ("labs.pdf", _make_text_pdf(SAMPLE_LAB_TEXT),
                                         "application/pdf")}).json()
    url = API + f"/documents/{result['id']}/reprocess"
    for _ in range(2):
        processed = client.post(url, headers=headers)
        assert processed.status_code == 200
        assert len(processed.json()["lab_results"]) == 3
    document = db_session.get(Document, result["id"])
    document.status = ProcessingStatus.PROCESSING
    db_session.commit()
    assert client.post(url, headers=headers).status_code == 409
