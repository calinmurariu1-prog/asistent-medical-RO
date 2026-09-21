"""Module 13 - full medical report export (PDF / Word).

`gather_report_data` builds a plain dict; `render_pdf` and `render_docx` turn it
into bytes. reportlab / python-docx are imported lazily so importing this module
never requires them (e.g. in unrelated tests).
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.appointment import Appointment
from app.models.clinical import MedicalHistory
from app.models.document import LabResult
from app.models.medication import Medication
from app.models.patient import Allergy, Patient, Vaccine
from app.models.user import User
from app.services import recommendations
from app.services.ai.base import DISCLAIMER

_TITLE = "Raport Medical - Asistent Medical AI"


_LABELS = {
    "male": "masculin",
    "female": "feminin",
    "other": "altul",
    "unspecified": "neprecizat",
    "unknown": "necunoscut",
    "observation": "observație",
    "diagnosis": "diagnostic",
    "procedure": "procedură",
    "surgery": "intervenție",
    "hospitalization": "internare",
    "vaccine": "vaccinare",
    "treatment": "tratament",
    "chronic_condition": "afecțiune cronică",
    "family_history": "istoric familial",
    "normal": "în interval",
    "high": "peste interval",
    "low": "sub interval",
    "critical_high": "marcat critic ridicat",
    "critical_low": "marcat critic scăzut",
    "verified": "parser / manual",
    "unverified": "de confirmat",
    "scheduled": "programată",
    "completed": "efectuată",
    "cancelled": "anulată",
    "missed": "neefectuată",
    "mild": "ușoară",
    "moderate": "moderată",
    "severe": "severă",
}


def _label(value: str) -> str:
    return _LABELS.get(value, value)


def gather_report_data(db: Session, patient: Patient, user: User) -> dict:
    bmi = None
    if patient.weight_kg and patient.height_cm:
        h = patient.height_cm / 100
        bmi = round(patient.weight_kg / (h * h), 1)

    history = db.scalars(
        select(MedicalHistory)
        .where(MedicalHistory.patient_id == patient.id)
        .order_by(MedicalHistory.event_date.desc().nullslast())
    ).all()

    meds = db.scalars(
        select(Medication).where(
            Medication.patient_id == patient.id, Medication.is_active.is_(True)
        )
    ).all()

    labs = db.scalars(
        select(LabResult)
        .where(LabResult.patient_id == patient.id)
        .order_by(LabResult.measured_on.desc().nullslast(), LabResult.id.desc())
    ).all()
    allergies = db.scalars(select(Allergy).where(Allergy.patient_id == patient.id)).all()
    vaccines = db.scalars(select(Vaccine).where(Vaccine.patient_id == patient.id)).all()
    appointments = db.scalars(
        select(Appointment)
        .where(Appointment.patient_id == patient.id)
        .order_by(Appointment.starts_at.desc())
    ).all()
    rec = recommendations.build_recommendations(db, patient)

    full_name = " ".join(p for p in [patient.first_name, patient.last_name] if p) or (
        user.full_name or user.email
    )

    return {
        "generated_at": datetime.now(UTC).strftime("%d.%m.%Y %H:%M"),
        "patient": {
            "name": full_name,
            "birth_date": patient.birth_date.isoformat() if patient.birth_date else "-",
            "sex": patient.sex.value if patient.sex else "-",
            "blood_type": patient.blood_type.value if patient.blood_type else "-",
            "weight_kg": patient.weight_kg,
            "height_cm": patient.height_cm,
            "bmi": bmi,
        },
        "history": [
            {
                "date": h.event_date.isoformat() if h.event_date else "-",
                "type": h.event_type.value,
                "title": h.title,
                "description": h.description or "",
            }
            for h in history
        ],
        "labs": [
            {
                "analyte": lab.analyte,
                "value": lab.value if lab.value is not None else (lab.value_text or "-"),
                "unit": lab.unit or "",
                "flag": lab.flag.value,
                "date": lab.measured_on.isoformat() if lab.measured_on else "Fără dată",
                "reference": f"{lab.ref_low} - {lab.ref_high}"
                if lab.ref_low is not None and lab.ref_high is not None
                else "Necunoscut",
                "confidence": lab.confidence,
            }
            for lab in labs
        ],
        "allergies": [
            {"substance": a.substance, "reaction": a.reaction or "-", "severity": a.severity.value}
            for a in allergies
        ],
        "vaccines": [
            {
                "name": v.name,
                "dose": v.dose or "-",
                "provider": v.provider or "-",
                "date": v.administered_on.isoformat() if v.administered_on else "-",
            }
            for v in vaccines
        ],
        "appointments": [
            {
                "title": a.title,
                "date": a.starts_at.isoformat(),
                "status": a.status.value,
                "location": a.location or "-",
                "notes": a.notes or "",
            }
            for a in appointments
        ],
        "medications": [
            {"name": m.name, "dose": m.dose or "", "frequency": m.frequency or ""} for m in meds
        ],
        "recommendations": {
            "alerts": rec.alerts,
            "questions_for_doctor": rec.questions_for_doctor,
        },
    }


def render_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        CondPageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=_TITLE)
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "MedicalDejaVu"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont(
                font_name, str(Path(__file__).resolve().parents[1] / "assets/fonts/DejaVuSans.ttf")
            )
        )
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font_name
    styles["Normal"].fontSize = 9
    styles["Normal"].leading = 13
    styles["Heading2"].keepWithNext = True
    story = []

    story.append(Paragraph(_TITLE, styles["Title"]))
    story.append(Paragraph(f"Generat: {data['generated_at']}", styles["Normal"]))
    story.append(Spacer(1, 0.4 * cm))

    p = data["patient"]
    story.append(Paragraph("Pacient", styles["Heading2"]))
    story.append(Paragraph(f"Nume: {escape(str(p['name']))}", styles["Normal"]))
    story.append(
        Paragraph(
            f"Naștere: {p['birth_date']} | Sex: {_label(p['sex'])} | Grup sanguin: "
            f"{_label(p['blood_type'])} | IMC: {p['bmi'] or '-'}",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 0.3 * cm))

    def _table(title: str, header: list[str], rows: list[list], empty: str) -> None:
        story.append(CondPageBreak(100))
        story.append(Paragraph(title, styles["Heading2"]))
        if not rows:
            story.append(Paragraph(empty, styles["Italic"]))
            story.append(Spacer(1, 0.3 * cm))
            return
        cells = [
            [
                Paragraph(escape(str(value if value is not None else "-")), styles["Normal"])
                for value in row
            ]
            for row in [header, *rows]
        ]
        table = Table(
            cells,
            colWidths=(
                [doc.width * ratio for ratio in (0.2, 0.2, 0.6)]
                if title == "Istoric medical"
                else [doc.width / len(header)] * len(header)
            ),
            hAlign="LEFT",
            repeatRows=1,
            splitInRow=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0E7FF")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 0.3 * cm))

    _table(
        "Istoric medical",
        ["Data", "Tip", "Descriere"],
        [
            [h["date"], _label(h["type"]), f"{h['title']}\n{h.get('description', '')}"]
            for h in data["history"]
        ],
        "Fără intrări.",
    )
    _table(
        "Analize (istoric complet)",
        ["Analit / data", "Valoare / unitate", "Referință", "Status / verificare"],
        [
            [
                f"{lr['analyte']} / {lr.get('date', '-')}",
                f"{lr['value']} {lr['unit']}",
                lr.get("reference", "-"),
                f"{_label(lr['flag'])} / {_label(lr.get('confidence', '-'))}",
            ]
            for lr in data["labs"]
        ],
        "Fără analize.",
    )
    _table(
        "Tratamente active",
        ["Medicament", "Doză", "Frecvență"],
        [[m["name"], m["dose"], m["frequency"]] for m in data["medications"]],
        "Fără tratamente active.",
    )

    _table(
        "Alergii",
        ["Substanță", "Reacție", "Severitate"],
        [[a["substance"], a["reaction"], _label(a["severity"])] for a in data.get("allergies", [])],
        "Nu sunt înregistrate alergii; aceasta nu confirmă absența lor.",
    )
    _table(
        "Vaccinări",
        ["Vaccin", "Doză", "Data", "Furnizor"],
        [[v["name"], v["dose"], v["date"], v["provider"]] for v in data.get("vaccines", [])],
        "Nu sunt înregistrate vaccinări.",
    )
    _table(
        "Programări",
        ["Consultație", "Data / status", "Locație / note"],
        [
            [a["title"], f"{a['date']} / {_label(a['status'])}", f"{a['location']} / {a['notes']}"]
            for a in data.get("appointments", [])
        ],
        "Nu sunt înregistrate programări.",
    )

    rec = data["recommendations"]
    story.append(Paragraph("Recomandări (orientative)", styles["Heading2"]))
    for line in rec["alerts"] + rec["questions_for_doctor"]:
        story.append(Paragraph(f"• {escape(str(line))}", styles["Normal"]))
    if not (rec["alerts"] or rec["questions_for_doctor"]):
        story.append(Paragraph("Fără recomandări.", styles["Italic"]))

    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph(DISCLAIMER, styles["Italic"]))

    def page_number(canvas, document):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Pagina {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return buf.getvalue()


def render_docx(data: dict) -> bytes:
    from docx import Document as Docx

    document = Docx()
    document.add_heading(_TITLE, level=0)
    document.add_paragraph(f"Generat: {data['generated_at']}")

    p = data["patient"]
    document.add_heading("Pacient", level=1)
    document.add_paragraph(f"Nume: {p['name']}")
    document.add_paragraph(
        f"Naștere: {p['birth_date']} | Sex: {_label(p['sex'])} | "
        f"Grup sanguin: {_label(p['blood_type'])} | IMC: {p['bmi'] or '-'}"
    )

    def _table(title: str, header: list[str], rows: list[list], empty: str) -> None:
        document.add_heading(title, level=1)
        if not rows:
            document.add_paragraph(empty)
            return
        table = document.add_table(rows=1, cols=len(header))
        table.style = "Light Grid Accent 1"
        for i, col in enumerate(header):
            table.rows[0].cells[i].text = col
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)

    _table(
        "Istoric medical",
        ["Data", "Tip", "Descriere"],
        [
            [h["date"], _label(h["type"]), f"{h['title']}\n{h.get('description', '')}"]
            for h in data["history"]
        ],
        "Fără intrări.",
    )
    _table(
        "Analize (istoric complet)",
        ["Analit / data", "Valoare / unitate", "Referință", "Status / verificare"],
        [
            [
                f"{lr['analyte']} / {lr.get('date', '-')}",
                f"{lr['value']} {lr['unit']}",
                lr.get("reference", "-"),
                f"{_label(lr['flag'])} / {_label(lr.get('confidence', '-'))}",
            ]
            for lr in data["labs"]
        ],
        "Fără analize.",
    )
    _table(
        "Tratamente active",
        ["Medicament", "Doză", "Frecvență"],
        [[m["name"], m["dose"], m["frequency"]] for m in data["medications"]],
        "Fără tratamente active.",
    )

    _table(
        "Alergii",
        ["Substanță", "Reacție", "Severitate"],
        [[a["substance"], a["reaction"], _label(a["severity"])] for a in data.get("allergies", [])],
        "Nu sunt înregistrate alergii; aceasta nu confirmă absența lor.",
    )
    _table(
        "Vaccinări",
        ["Vaccin", "Doză", "Data", "Furnizor"],
        [[v["name"], v["dose"], v["date"], v["provider"]] for v in data.get("vaccines", [])],
        "Nu sunt înregistrate vaccinări.",
    )
    _table(
        "Programări",
        ["Consultație", "Data / status", "Locație / note"],
        [
            [a["title"], f"{a['date']} / {_label(a['status'])}", f"{a['location']} / {a['notes']}"]
            for a in data.get("appointments", [])
        ],
        "Nu sunt înregistrate programări.",
    )

    rec = data["recommendations"]
    document.add_heading("Recomandări (orientative)", level=1)
    lines = rec["alerts"] + rec["questions_for_doctor"]
    for line in lines:
        document.add_paragraph(line, style="List Bullet")
    if not lines:
        document.add_paragraph("Fără recomandări.")

    document.add_paragraph()
    document.add_paragraph(DISCLAIMER)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
