"""Export documente în PDF și Word (DOCX), cu suport pentru diacritice.

Pentru PDF se încearcă înregistrarea unei fonturi Unicode (DejaVuSans) ca
diacriticele românești (ș, ț, ă, â, î) să apară corect; dacă fontul nu e
găsit, se cade pe Helvetica.
"""
from __future__ import annotations

import io
import os

# --------------------------------------------------------------------------
# PDF (reportlab)
# --------------------------------------------------------------------------
_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
]


def _register_unicode_font() -> tuple[str, str]:
    """Returnează (font_normal, font_bold); cade pe Helvetica dacă lipsește."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular = next((p for p in _FONT_CANDIDATES if "Bold" not in p and os.path.exists(p)), None)
    bold = next((p for p in _FONT_CANDIDATES if "Bold" in p and os.path.exists(p)), None)
    if not regular:
        return "Helvetica", "Helvetica-Bold"
    try:
        pdfmetrics.registerFont(TTFont("Doc", regular))
        if bold:
            pdfmetrics.registerFont(TTFont("Doc-Bold", bold))
            return "Doc", "Doc-Bold"
        return "Doc", "Doc"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def to_pdf(title: str, content: str) -> bytes:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font, font_bold = _register_unicode_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=title,
    )
    title_style = ParagraphStyle("Title", fontName=font_bold, fontSize=16, spaceAfter=14, leading=20)
    body_style = ParagraphStyle("Body", fontName=font, fontSize=11, leading=16, alignment=TA_LEFT)

    story = [Paragraph(_esc(title), title_style), Spacer(1, 6)]
    for line in content.split("\n"):
        story.append(Paragraph(_esc(line) if line.strip() else "&nbsp;", body_style))
    doc.build(story)
    return buf.getvalue()


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# --------------------------------------------------------------------------
# DOCX (python-docx)
# --------------------------------------------------------------------------
def to_docx(title: str, content: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    heading = document.add_heading(title, level=1)
    heading.runs[0].font.size = Pt(16)
    for line in content.split("\n"):
        document.add_paragraph(line)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
